#!/usr/bin/env python3
"""Convierte el export XLSX de Sellerboard al formato de ventas del dashboard.

Uso:
  python3 scripts/import_sellerboard.py export.xlsx --output-dir .

El resultado es ventas_sellerboard_YYYY.csv. La importación es idempotente:
el fichero se reemplaza solo después de validar todas las filas.
"""
from __future__ import annotations

import argparse
import csv
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

FIELDS = [
    "amazon-order-id", "purchase-date", "sku", "quantity", "item-price",
    "sales-channel", "order-status", "sellerboard-sales", "sellerboard-refund-cost",
]


def money(value) -> str:
    if value in (None, "", "-"):
        return "0"
    return str(float(str(value).replace(",", ".")))


def refund_money(value) -> str:
    return str(abs(float(money(value))))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--output-dir", type=Path, default=Path("."))
    args = ap.parse_args()

    ws = load_workbook(args.xlsx, read_only=True, data_only=True).active
    headers = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True))]
    aliases = {h.lower().replace("с", "c"): h for h in headers}
    required = {"order number", "order date", "sku", "units", "sales", "refund cost"}
    missing = required.difference(aliases)
    if missing:
        raise SystemExit(f"Export Sellerboard incompatible; faltan: {', '.join(sorted(missing))}")
    pos = {name: headers.index(source) for name, source in aliases.items() if name in required}
    rows_by_year: dict[int, list[dict[str, str]]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        raw_date = str(values[pos["order date"]]).strip()
        date = datetime.strptime(raw_date, "%d.%m.%Y").strftime("%Y-%m-%dT00:00:00Z")
        order = str(values[pos["order number"]] or "").strip()
        status = "Returned" if "/ Return /" in order else ("Unshipped" if "/ Unshipped /" in order else "Shipped")
        channel = "Amazon"
        row = {
            "amazon-order-id": order.split(" /")[0],
            "purchase-date": date,
            "sku": str(values[pos["sku"]] or "").strip(),
            "quantity": str(values[pos["units"]] or "1"),
            "item-price": money(values[pos["sales"]]),
            "sales-channel": channel,
            "order-status": status,
            "sellerboard-sales": money(values[pos["sales"]]),
            "sellerboard-refund-cost": refund_money(values[pos["refund cost"]]),
        }
        rows_by_year.setdefault(datetime.strptime(raw_date, "%d.%m.%Y").year, []).append(row)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    for year, rows in rows_by_year.items():
        target = args.output_dir / f"ventas_sellerboard_{year}.csv"
        fd, tmp = tempfile.mkstemp(prefix=target.name + ".", dir=args.output_dir, text=True)
        try:
            with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=FIELDS, delimiter=";")
                writer.writeheader()
                writer.writerows(rows)
            os.replace(tmp, target)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
        print(f"{target}: {len(rows)} filas")


if __name__ == "__main__":
    main()
