import json, re, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
S=sys.argv[1]; OUT=sys.argv[2]
D=json.load(open(f"{S}/groups.json")); G=D['groups']
L=[json.loads(l) for l in open(f"{S}/listings_cache.jsonl")]
WRONG_PT={'tetera':'TEAPOT','lampara':'LAMP','bolsa':'GIFT_BAG / PAPER_BAG','bandeja':'TRAY','cuenco':'DISHWARE_BOWL','plato':'DISHWARE_PLATE','especiero':'SPICE_RACK','jarron':'VASE','portafotos':'PICTURE_FRAME','hucha':'MONEY_BANK','tabla':'CUTTING_BOARD','estanteria':'SHELF','reloj':'CLOCK','cesta':'BASKET','alfombra':'RUG','papel':'GIFT_WRAP','caja':'STORAGE_BOX','espejo':'MIRROR','vela':'CANDLE_HOLDER','portavelas':'CANDLE_HOLDER','lienzo':'WALL_ART','cuadro':'WALL_ART','taza':'CUP','neceser':'COSMETIC_CASE'}
def cap(s): return ' '.join(w if w.isupper() else w.capitalize() for w in s.split())
def parent_title(x):
    t=min((d['name'] for d in x['items']),key=len)
    t=re.sub(r'^(ROCKING GIFTS)\s*','',t,flags=re.I)
    t=re.split(r'\s[|,(]\s?|\s-\s',t)[0]
    t=re.sub(r'\bx\s?\d+(\s+modelos)?\b','',t,flags=re.I)
    t=re.sub(r'\d+([.,]\d+)?(\s*x\s*\d+([.,]\d+)?)*\s*(cm|mm|ml|l|cl)\b','',t,flags=re.I)
    t=re.sub(r'\b\d+([.,]\d+)?\s*l\b','',t,flags=re.I)
    from norm import COLORS
    t=' '.join(w for w in t.split() if w.lower().strip(',') not in COLORS)
    t=re.sub(r'\b\d+\s+(cajones|niveles|cuencos|departamentos|brazos|velas)\b','',t,flags=re.I)
    t=re.sub(r'\s+',' ',t).strip(' -,')
    t=re.sub(r'(\s+(y|de|con|en|del))+$','',t,flags=re.I)
    dims={'SIZE_NAME':'Varios Tamaños','COLOR_NAME':'Varios Colores','STYLE_NAME':'Varios Diseños','NUMBER_OF_ITEMS':'Varios Packs'}
    suf=' y '.join(dims[p] for p in x['theme'].replace('SIZE/','SIZE_NAME/').split('/') if p in dims)
    return f"ROCKING GIFTS {t} - {suf}" if suf else f"ROCKING GIFTS {t}"
rows=[]
keypt=collections.Counter(x['key'] for x in G)
ERR={x['sku'] for x in L if any(i.endswith(':ERROR') for i in (x.get('issues') or []))}
for i,x in enumerate(G,1):
    x['id']=f"VAR-{i:03d}"
    n=x['n']; act=x['active']
    warn=[]
    if x['need_style']: warn.append('Definir STYLE_NAME (nombre de diseño) por imagen: hay hijos con mismo tamaño/color')
    if x['packvar']: warn.append('Mezcla unidades sueltas y packs: indicar pack en SIZE_NAME/STYLE_NAME ("x2")')
    if x['pmin'] and x['pmax']/x['pmin']>6: warn.append(f"Rango de precio amplio ({x['pmin']}-{x['pmax']} €): revisar que sea el mismo producto")
    k0=x['key'].split()
    sug=[v for kk,v in WRONG_PT.items() if kk in k0]
    if sug and x['pt'] in ('SCULPTURE','FIGURINE','HOME','HOME_FURNITURE_AND_DECOR'): warn.append(f"productType {x['pt']} poco preciso (sería {sug[0]}); se mantiene para no romper la familia")
    if 'letra' in k0: warn.append('Usar theme LETTER_CHARACTER (COLOR/LETTER_CHARACTER) en vez de STYLE_NAME'); x['theme']='COLOR/LETTER_CHARACTER'
    if keypt[x['key']]>1: warn.append('Existe grupo gemelo con otro productType: se podrían unir si se unifica productType (requiere caso a Soporte)')
    ne=sum(1 for d in x['items'] if d['sku'] in ERR)
    if ne: warn.append(f'{ne} hijo(s) con ERROR en Amazon: corregir antes')
    if act<n: warn.append(f"{n-act} hijo(s) inactivos: activar o excluir antes de crear")
    x['warn']=warn
    score=n*2+act+min(x['units'],20)
    x['prio']='A' if (n>=5 and act>=3) or (n>=4 and act==n) else ('B' if n>=3 or act==n else 'C')
    x['score']=score; x['ptitle']=parent_title(x)
G.sort(key=lambda x:({'A':0,'B':1,'C':2}[x['prio']],-x['score']))
wb=Workbook(); H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='232F3E')
def sheet(ws,hdr,data,widths):
    ws.append(hdr)
    for c in ws[1]: c.font=H; c.fill=F
    for r in data: ws.append(r)
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Grupos'
sheet(ws,['ID','Prioridad','Título padre propuesto','productType','variation_theme','Nº hijos','Activos','Precio mín','Precio máx','Proveedor','Uds vendidas 25-26','Avisos'],
 [[x['id'],x['prio'],x['ptitle'],x['pt'],x['theme'],x['n'],x['active'],x['pmin'],x['pmax'],{'SG':'Signes Grimalt','DC':'Dcasa','DCSG':'Mixto'}[x['suppliers']],x['units'],' | '.join(x['warn'])] for x in G],
 [10,9,70,22,24,9,8,10,10,14,12,90])
ws2=wb.create_sheet('Hijos')
data=[]
for x in G:
    for d in sorted(x['items'],key=lambda d:float(d['price'] or 0)):
        size=(d['dim'] or d['size']).replace('x',' x ').replace('cm',' cm').strip()
        if x['packvar'] and d['pack']>1: size+=f" (x{d['pack']})"
        data.append([x['id'],x['prio'],d['sku'],d['asin'],d['name'],d['price'],d['status'],size if 'SIZE' in x['theme'] else '',
            (d['color'] or d['colort'].title()) if 'COLOR' in x['theme'] else '', 'DEFINIR' if 'STYLE' in x['theme'] else '', d['pack'], d.get('image') or ''])
sheet(ws2,['ID grupo','Prioridad','SKU','ASIN','Título actual','Precio','Estado','size_name','color_name','style_name','Unidades pack','Imagen'],data,[10,9,14,13,80,8,10,20,16,14,8,50])
# auditoría
ws3=wb.create_sheet('Auditoría')
brands=collections.Counter(((x.get('attrs') or {}).get('brand') or [{'value':'(sin contenido propio)'}])[0]['value'] for x in L)
issues=collections.Counter(i for x in L for i in (x.get('issues') or []) )
ptc=collections.Counter(x.get('productType') for x in L)
aud=[['Total SKUs Signes Grimalt + Dcasa analizados',len(L)],['Elegibles ROCKING GIFTS (sin variación previa)',D['total_rg']]]
aud+=[['Excluido: '+k,v] for k,v in D['excl'].items()]
aud+=[['Marca: '+k,v] for k,v in brands.most_common()]
aud+=[['SKUs con issues de Amazon',sum(1 for x in L if x.get('issues'))]]
aud+=[['Issue '+k,v] for k,v in issues.most_common(15)]
aud+=[['productType '+str(k),v] for k,v in ptc.most_common(25)]
aud+=[['Grupos propuestos',len(G)],['SKUs en grupos',sum(x['n'] for x in G)]]+[[f'Prioridad {p}',sum(1 for x in G if x['prio']==p)] for p in 'ABC']
sheet(ws3,['Métrica','Valor'],aud,[60,12])
wb.save(OUT)
json.dump(G,open(f"{S}/groups_final.json","w"),ensure_ascii=False,default=str)
print(len(G),collections.Counter(x['prio'] for x in G),sum(x['n'] for x in G))
print(collections.Counter(x['theme'] for x in G))
print('issues',issues.most_common(8)); print('brands',brands.most_common(6)); print('pt',ptc.most_common(12))
for x in G[:25]: print(x['id'],x['prio'],x['n'],x['theme'],'|',x['ptitle'])
